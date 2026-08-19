"""Build the SRE Agent executive demo deck as a real .pptx file."""
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

# ---- palette (matches the HTML artifact) ----
BG = RGBColor(0xFA, 0xFA, 0xF7)
INK = RGBColor(0x26, 0x28, 0x2B)
INK_SOFT = RGBColor(0x56, 0x5A, 0x5E)
INK_FAINT = RGBColor(0x8A, 0x8E, 0x92)
BLUE = RGBColor(0x4A, 0x7C, 0xA8)
BLUE_BG = RGBColor(0xEA, 0xF1, 0xF7)
BLUE_LINE = RGBColor(0xC9, 0xDC, 0xEA)
TEAL = RGBColor(0x3E, 0x8F, 0x7C)
TEAL_BG = RGBColor(0xE7, 0xF3, 0xEF)
TEAL_LINE = RGBColor(0xC4, 0xE1, 0xD8)
AMBER = RGBColor(0xB8, 0x86, 0x3C)
AMBER_BG = RGBColor(0xF7, 0xEF, 0xDF)
AMBER_LINE = RGBColor(0xE6, 0xD2, 0xAA)
GRAY_BG = RGBColor(0xF1, 0xF0, 0xEC)
CARD_BORDER = RGBColor(0xE4, 0xE1, 0xD8)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

SERIF = "Georgia"
SANS = "Calibri"

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]
SW, SH = prs.slide_width, prs.slide_height


def add_bg(slide, color=BG):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = color


def txt(slide, x, y, w, h, text, size=14, color=INK, bold=False, italic=False,
         font=SANS, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, line_spacing=1.0,
         wrap=True):
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = wrap
    tf.vertical_anchor = anchor
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    lines = text.split("\n")
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = line_spacing
        r = p.add_run()
        r.text = line
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.italic = italic
        r.font.name = font
        r.font.color.rgb = color
    return box


def rounded_card(slide, x, y, w, h, fill=WHITE, line=CARD_BORDER, radius=0.08):
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    try:
        shp.adjustments[0] = radius
    except Exception:
        pass
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    shp.line.color.rgb = line
    shp.line.width = Pt(0.75)
    shp.shadow.inherit = False
    return shp


def eyebrow(slide, text, x=Inches(0.6), y=Inches(0.42)):
    w = Inches(0.14 * len(text) + 0.5)
    pill = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, Inches(0.36))
    pill.adjustments[0] = 0.5
    pill.fill.solid()
    pill.fill.fore_color.rgb = BLUE_BG
    pill.line.color.rgb = BLUE_LINE
    pill.line.width = Pt(0.75)
    pill.shadow.inherit = False
    tf = pill.text_frame
    tf.margin_left = Inches(0.08); tf.margin_right = Inches(0.08)
    tf.margin_top = 0; tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = text.upper()
    r.font.size = Pt(10.5)
    r.font.bold = True
    r.font.name = SANS
    r.font.color.rgb = BLUE


def title(slide, text, x=Inches(0.6), y=Inches(0.9), w=Inches(10.5), size=32):
    txt(slide, x, y, w, Inches(1.1), text, size=size, color=INK, font=SERIF,
        line_spacing=1.05)


def set_notes(slide, notes):
    slide.notes_slide.notes_text_frame.text = notes


def checklist(slide, x, y, w, items, accent=TEAL, text_color=INK, size=13, dot=True):
    cy = y
    for item in items:
        mark = slide.shapes.add_shape(MSO_SHAPE.OVAL, x, cy + Inches(0.04), Inches(0.12), Inches(0.12))
        mark.fill.solid()
        mark.fill.fore_color.rgb = accent
        mark.line.fill.background()
        mark.shadow.inherit = False
        txt(slide, x + Inches(0.24), cy, w - Inches(0.24), Inches(0.32), item,
            size=size, color=text_color, font=SANS)
        cy += Inches(0.36)
    return cy


def flow_row(slide, y, steps, w_total=Inches(11.8), x0=Inches(0.75), h=Inches(1.05)):
    n = len(steps)
    gap = Inches(0.35)
    step_w = Emu(int((w_total - gap * (n - 1)) / n))
    x = x0
    for i, (label, detail) in enumerate(steps):
        card = rounded_card(slide, x, y, step_w, h, fill=WHITE, line=CARD_BORDER, radius=0.12)
        tf = card.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf.margin_left = Inches(0.08); tf.margin_right = Inches(0.08)
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run(); r.text = label
        r.font.size = Pt(13); r.font.bold = True; r.font.name = SANS; r.font.color.rgb = INK
        if detail:
            p2 = tf.add_paragraph()
            p2.alignment = PP_ALIGN.CENTER
            r2 = p2.add_run(); r2.text = detail
            r2.font.size = Pt(9.5); r2.font.name = SANS; r2.font.color.rgb = INK_FAINT
        if i < n - 1:
            ax = Emu(int(x + step_w + gap / 2))
            arrow = slide.shapes.add_textbox(Emu(int(ax - Inches(0.15))), y, Inches(0.3), h)
            atf = arrow.text_frame
            atf.vertical_anchor = MSO_ANCHOR.MIDDLE
            ap = atf.paragraphs[0]; ap.alignment = PP_ALIGN.CENTER
            ar = ap.add_run(); ar.text = "→"
            ar.font.size = Pt(16); ar.font.color.rgb = INK_FAINT
        x = Emu(int(x + step_w + gap))


# =========================================================
# SLIDE 1
# =========================================================
s = prs.slides.add_slide(BLANK); add_bg(s)
eyebrow(s, "Why should I care?")
title(s, "From Incident Alert to\nEvidence-Backed RCA", y=Inches(0.85), size=34)
txt(s, Inches(0.6), Inches(2.05), Inches(9.8), Inches(0.7),
    "Production incidents often require engineers to manually correlate information across "
    "multiple operational systems before they can confidently identify what happened.",
    size=15, color=INK_SOFT, line_spacing=1.2)

flow_row(s, Inches(2.75), [
    ("Incident occurs", ""),
    ("Engineer investigates", "multiple signals, manually"),
    ("Root cause identified", ""),
])

note_card = rounded_card(s, Inches(0.75), Inches(4.15), Inches(11.8), Inches(0.85), fill=TEAL_BG, line=TEAL_LINE, radius=0.15)
tf = note_card.text_frame; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
tf.margin_left = Inches(0.35); tf.margin_right = Inches(0.35)
p = tf.paragraphs[0]
r1 = p.add_run(); r1.text = "The opportunity:  "; r1.font.bold = True; r1.font.size = Pt(14); r1.font.color.rgb = INK; r1.font.name = SANS
r2 = p.add_run(); r2.text = "SRE Agent investigates the evidence automatically and gives the engineer a structured starting point."
r2.font.size = Pt(14); r2.font.color.rgb = INK; r2.font.name = SANS

bullets1 = ["Reduce investigation effort", "Gather evidence consistently",
            "Accelerate root-cause understanding", "Keep engineers in control"]
bx = Inches(0.75)
for b in bullets1:
    checklist(s, bx, Inches(5.35), Inches(2.7), [b], accent=TEAL, size=12.5)
    bx = Emu(int(bx + Inches(2.9)))

set_notes(s,
    "Today, when an incident occurs, an engineer still has to gather information from several "
    "places before they can understand what happened. The idea behind this project is "
    "straightforward: let the agent perform the initial investigation and evidence gathering "
    "automatically. The engineer receives an evidence-backed root-cause analysis instead of "
    "starting from an alert with very little context. The goal is not to replace the SRE. It is "
    "to give the SRE a much stronger starting point.")

# =========================================================
# SLIDE 2
# =========================================================
s = prs.slides.add_slide(BLANK); add_bg(s)
eyebrow(s, "Why is incident investigation difficult?")
title(s, "The Operational Problem", size=32)

flow_row(s, Inches(2.3), [
    ("Alert", "Limited initial context"),
    ("Investigation", "Logs, events, pod state,\nconfiguration, cluster context"),
    ("Engineering judgment", "Correlate evidence"),
    ("RCA", "Determine likely cause"),
], h=Inches(1.3))

line = s.shapes.add_connector(1, Inches(0.75), Inches(4.15), Inches(12.55), Inches(4.15))
line.line.color.rgb = CARD_BORDER
line.line.width = Pt(1)

txt(s, Inches(0.6), Inches(4.4), Inches(12.1), Inches(0.5),
    "The bottleneck is not alerting.", size=18, color=INK_FAINT, font=SERIF, italic=True, align=PP_ALIGN.CENTER)
txt(s, Inches(0.6), Inches(4.9), Inches(12.1), Inches(0.55),
    "The bottleneck is investigation and correlation.", size=20, color=BLUE, font=SERIF, italic=True, align=PP_ALIGN.CENTER)

set_notes(s,
    "We already have systems that tell us when something is wrong. The difficult part starts "
    "after the alert. An engineer has to determine which cluster is affected, inspect "
    "Kubernetes state, look at events and logs, review configuration, and then correlate those "
    "signals. That investigation requires time and experience. This is the part of the incident "
    "lifecycle we are targeting.")

# =========================================================
# SLIDE 3
# =========================================================
s = prs.slides.add_slide(BLANK); add_bg(s)
eyebrow(s, "What did we build?")
title(s, "An AI Investigator for Kubernetes Incidents", size=30, w=Inches(11.5))

cap_data = [
    ("Investigate", "Collect real operational evidence directly from the environment.", BLUE_BG, BLUE_LINE, BLUE),
    ("Reason", "Correlate signals and weigh competing hypotheses.", TEAL_BG, TEAL_LINE, TEAL),
    ("Explain", "Produce an evidence-backed RCA, not just an answer.", AMBER_BG, AMBER_LINE, AMBER),
]
cw = Inches(3.75); gap = Inches(0.3); x = Inches(0.75); y = Inches(2.15)
for label, desc, bgc, linec, accent in cap_data:
    card = rounded_card(s, x, y, cw, Inches(1.8), fill=WHITE, line=CARD_BORDER, radius=0.1)
    dot = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Emu(int(x + Inches(0.25))), Emu(int(y + Inches(0.25))), Inches(0.5), Inches(0.5))
    dot.adjustments[0] = 0.25
    dot.fill.solid(); dot.fill.fore_color.rgb = bgc
    dot.line.color.rgb = linec; dot.line.width = Pt(0.75)
    dot.shadow.inherit = False
    txt(s, Emu(int(x + Inches(0.25))), Emu(int(y + 0.95 * 914400)), cw - Inches(0.5), Inches(0.35),
        label, size=16, bold=True, color=INK, font=SANS)
    txt(s, Emu(int(x + Inches(0.25))), Emu(int(y + 1.3 * 914400)), cw - Inches(0.5), Inches(0.5),
        desc, size=11.5, color=INK_SOFT, font=SANS, line_spacing=1.15)
    x = Emu(int(x + cw + gap))

principles = ["Read-only by default", "Evidence before conclusions", "Human-controlled remediation"]
x = Inches(0.75); y = Inches(4.3)
for p in principles:
    w = Inches(0.11 * len(p) + 0.6)
    chip = rounded_card(s, x, y, w, Inches(0.45), fill=GRAY_BG, line=CARD_BORDER, radius=0.3)
    tf = chip.text_frame; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = Inches(0.15); tf.margin_right = Inches(0.15)
    pp = tf.paragraphs[0]; pp.alignment = PP_ALIGN.CENTER
    r = pp.add_run(); r.text = p
    r.font.size = Pt(12); r.font.bold = True; r.font.color.rgb = INK; r.font.name = SANS
    x = Emu(int(x + w + Inches(0.25)))

set_notes(s,
    "We built an SRE Agent focused specifically on investigation. It collects real operational "
    "evidence rather than answering only from the alert description. It reasons across those "
    "signals and produces an RCA explaining what it believes happened and why. Investigation "
    "remains read-only by default. The engineer still owns the remediation decision.")

# =========================================================
# SLIDE 4
# =========================================================
s = prs.slides.add_slide(BLANK); add_bg(s)
eyebrow(s, "How does this technically work?")
title(s, "How the Agent Works", size=30)

lanes = [
    ("Trigger", "01", "Incident", "PagerDuty, Engineer, API", WHITE, CARD_BORDER),
    ("Intelligence", "02", "SRE Agent", "Vertex AI Agent Engine\nInvestigation workflow", BLUE_BG, BLUE_LINE),
    ("Access", "03", "Tool Access", "Agent Gateway\nMCP integrations", WHITE, CARD_BORDER),
    ("Infrastructure", "04", "Environment", "GKE / Kubernetes\nFuture non-GKE targets", WHITE, CARD_BORDER),
    ("Evidence", "05", "Evidence", "Logs, events\nPod status, config", TEAL_BG, TEAL_LINE),
    ("Decision Support", "06", "Outcome", "Evidence-backed RCA\nConfidence, human review", WHITE, CARD_BORDER),
]
n = len(lanes)
w_total = Inches(11.8); gap = Inches(0.22); x0 = Inches(0.75)
lane_w = Emu(int((w_total - gap * (n - 1)) / n))
x = x0
laney = Inches(2.35)
for lbl, num, ttl, sub, bgc, linec in lanes:
    txt(s, x, Inches(2.0), lane_w, Inches(0.3), lbl.upper(), size=8.5, bold=True, color=INK_FAINT,
        align=PP_ALIGN.CENTER, font=SANS)
    card = rounded_card(s, x, laney, lane_w, Inches(1.35), fill=bgc, line=linec, radius=0.12)
    tf = card.text_frame; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.word_wrap = True
    tf.margin_left = Inches(0.06); tf.margin_right = Inches(0.06)
    p0 = tf.paragraphs[0]; p0.alignment = PP_ALIGN.CENTER
    r0 = p0.add_run(); r0.text = num
    r0.font.size = Pt(8.5); r0.font.bold = True; r0.font.color.rgb = INK_FAINT; r0.font.name = SANS
    p1 = tf.add_paragraph(); p1.alignment = PP_ALIGN.CENTER
    r1 = p1.add_run(); r1.text = ttl
    r1.font.size = Pt(12.5); r1.font.bold = True; r1.font.color.rgb = INK; r1.font.name = SANS
    p2 = tf.add_paragraph(); p2.alignment = PP_ALIGN.CENTER
    r2 = p2.add_run(); r2.text = sub
    r2.font.size = Pt(8.5); r2.font.color.rgb = INK_FAINT; r2.font.name = SANS
    if lbl != "Decision Support":
        ax = Emu(int(x + lane_w))
        arrow = s.shapes.add_textbox(ax, laney, gap, Inches(1.35))
        atf = arrow.text_frame; atf.vertical_anchor = MSO_ANCHOR.MIDDLE
        ap = atf.paragraphs[0]; ap.alignment = PP_ALIGN.CENTER
        ar = ap.add_run(); ar.text = "→"
        ar.font.size = Pt(13); ar.font.color.rgb = INK_FAINT
    x = Emu(int(x + lane_w + gap))

principles2 = ["Multi-cluster by design", "Read-only investigation", "Extensible data sources"]
x = Inches(0.75); y = Inches(4.35)
for p in principles2:
    w = Inches(0.11 * len(p) + 0.6)
    chip = rounded_card(s, x, y, w, Inches(0.45), fill=GRAY_BG, line=CARD_BORDER, radius=0.3)
    tf = chip.text_frame; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = Inches(0.15); tf.margin_right = Inches(0.15)
    pp = tf.paragraphs[0]; pp.alignment = PP_ALIGN.CENTER
    r = pp.add_run(); r.text = p
    r.font.size = Pt(12); r.font.bold = True; r.font.color.rgb = INK; r.font.name = SANS
    x = Emu(int(x + w + Inches(0.25)))

set_notes(s,
    "At a high level, the agent runs on Google Cloud using Vertex AI Agent Engine. When an "
    "incident arrives, the workflow determines where it needs to investigate. Through the Agent "
    "Gateway and MCP tools, it connects to the operational environment and collects real "
    "evidence. That includes Kubernetes events, logs, pod state and relevant configuration. The "
    "agent then reasons over the evidence and produces an RCA for the engineer. The architecture "
    "is intentionally designed so we can add clusters and additional data sources without "
    "changing the core investigation model.")

# =========================================================
# SLIDE 5
# =========================================================
s = prs.slides.add_slide(BLANK); add_bg(s)
eyebrow(s, "What actually works today?")
title(s, "Current State", size=30)

col_w = Inches(5.75)
left_card = rounded_card(s, Inches(0.75), Inches(2.0), col_w, Inches(4.7), fill=TEAL_BG, line=TEAL_LINE, radius=0.06)
right_card = rounded_card(s, Inches(6.75), Inches(2.0), col_w, Inches(4.7), fill=GRAY_BG, line=CARD_BORDER, radius=0.06)

txt(s, Inches(1.05), Inches(2.25), col_w - Inches(0.6), Inches(0.35), "VALIDATED TODAY",
    size=13, bold=True, color=TEAL, font=SANS)
checklist(s, Inches(1.05), Inches(2.7), col_w - Inches(0.6), [
    "Real Kubernetes investigations, live-verified",
    "GKE evidence collection (logs, events, pod state)",
    "Multi-cluster routing and support",
    "MCP tool execution against live infrastructure",
    "Evidence-backed RCA with confidence scoring",
    "Read-only investigation model",
    "Human-controlled remediation",
], accent=TEAL, size=12.5)

txt(s, Inches(7.05), Inches(2.25), col_w - Inches(0.6), Inches(0.35), "IN PROGRESS / NEXT",
    size=13, bold=True, color=INK_SOFT, font=SANS)
checklist(s, Inches(7.05), Inches(2.7), col_w - Inches(0.6), [
    "PagerDuty-triggered investigations",
    "Non-GKE cluster support",
    "Confidence calibration & improvements",
    "Expanded evaluation coverage",
    "Production monitoring & dashboards",
    "Additional incident scenario coverage",
], accent=INK_FAINT, size=12.5)

set_notes(s,
    "This is not a future concept deck. A meaningful portion of the investigation workflow "
    "already works today. We have validated real Kubernetes evidence collection, multi-cluster "
    "investigation and RCA generation. There are also areas we are deliberately still validating "
    "before calling this production-ready. I want to be very clear about that distinction because "
    "passing CI alone is not our definition of success.")

# =========================================================
# SLIDE 6
# =========================================================
s = prs.slides.add_slide(BLANK); add_bg(s)
eyebrow(s, "Can you prove it?")
title(s, "Live Demonstration", size=30)

steps = ["Introduce\nincident", "Agent\ninvestigates", "Evidence is\ncollected", "Agent forms\nRCA", "Engineer reviews\noutcome"]
n = len(steps)
sw_ = Inches(2.1); gap = Inches(0.15)
total = sw_ * n + gap * (n - 1)
x0 = Emu(int((SW - total) / 2))
x = x0
colors = [BLUE, BLUE, TEAL, TEAL, AMBER]
for i, (lbl, accent) in enumerate(zip(steps, colors)):
    cx = Emu(int(x + sw_ / 2))
    circ = s.shapes.add_shape(MSO_SHAPE.OVAL, Emu(int(cx - Inches(0.4))), Inches(2.3), Inches(0.8), Inches(0.8))
    circ.fill.solid(); circ.fill.fore_color.rgb = WHITE
    circ.line.color.rgb = accent; circ.line.width = Pt(1.5)
    circ.shadow.inherit = False
    tf = circ.text_frame; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = str(i + 1)
    r.font.size = Pt(20); r.font.name = SERIF; r.font.color.rgb = accent
    txt(s, x, Inches(3.25), sw_, Inches(0.6), lbl, size=12.5, bold=True, color=INK,
        align=PP_ALIGN.CENTER, font=SANS)
    x = Emu(int(x + sw_ + gap))

watch = rounded_card(s, Inches(1.9), Inches(4.3), Inches(9.5), Inches(1.1), fill=WHITE, line=CARD_BORDER, radius=0.1)
watch.line.color.rgb = BLUE
tf = watch.text_frame; tf.word_wrap = True; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
tf.margin_left = Inches(0.3); tf.margin_right = Inches(0.3)
p = tf.paragraphs[0]
r1 = p.add_run(); r1.text = "What to watch for:  "; r1.font.bold = True; r1.font.size = Pt(13.5); r1.font.color.rgb = INK; r1.font.name = SANS
r2 = p.add_run(); r2.text = ("not simply whether the agent produces an answer — whether its "
                              "conclusion is actually supported by the evidence it collected along the way.")
r2.font.size = Pt(13.5); r2.font.color.rgb = INK_SOFT; r2.font.name = SANS

set_notes(s,
    "Rather than spend more time describing it, I want to show the workflow. I'll introduce a "
    "real Kubernetes failure scenario. We'll watch the agent investigate it, see the evidence it "
    "collects, and then review the RCA it returns. The important thing to watch is not simply "
    "whether it produces an answer. Watch whether the conclusion is supported by the evidence it "
    "actually collected.")

# =========================================================
# SLIDE 7
# =========================================================
s = prs.slides.add_slide(BLANK); add_bg(s)
eyebrow(s, "Can this become more valuable at scale?")
title(s, "Where This Goes Next", size=30)

cols = [
    ("TODAY", TEAL, TEAL_LINE, ["Kubernetes investigation", "Multi-cluster support",
     "Evidence-backed RCA", "Human-controlled actions"]),
    ("NEXT", BLUE, BLUE_LINE, ["PagerDuty-triggered investigations", "Non-GKE environments",
     "Stronger confidence model", "Golden incident evaluation", "Production observability"]),
    ("SCALE", AMBER, AMBER_LINE, ["Additional operational data sources", "Cross-platform investigations",
     "Continuous evaluation", "Enterprise governance", "Broader incident intelligence"]),
]
cw = Inches(3.75); gap = Inches(0.35); x = Inches(0.75); y = Inches(2.15)
for label, accent, accent_line, items in cols:
    txt(s, x, y, cw, Inches(0.35), label, size=13, bold=True, color=accent, font=SANS)
    ln = s.shapes.add_connector(1, x, Emu(int(y + Inches(0.42))), Emu(int(x + cw)), Emu(int(y + Inches(0.42))))
    ln.line.color.rgb = accent_line
    ln.line.width = Pt(2)
    checklist(s, x, Emu(int(y + Inches(0.6))), cw, items, accent=accent, size=12, dot=True)
    if label != "SCALE":
        ax = Emu(int(x + cw + gap / 2))
        arrow = s.shapes.add_textbox(Emu(int(ax - Inches(0.15))), Emu(int(y + Inches(1.4))), Inches(0.3), Inches(0.4))
        atf = arrow.text_frame; atf.vertical_anchor = MSO_ANCHOR.MIDDLE
        ap = atf.paragraphs[0]; ap.alignment = PP_ALIGN.CENTER
        ar = ap.add_run(); ar.text = "→"
        ar.font.size = Pt(16); ar.font.color.rgb = INK_FAINT
    x = Emu(int(x + cw + gap))

set_notes(s,
    "The current implementation establishes the investigation foundation. The next step is "
    "connecting it more naturally into the incident lifecycle and strengthening how we measure "
    "investigation quality. The architecture also gives us a path beyond one Kubernetes "
    "environment. Over time, the same investigation model can incorporate additional "
    "infrastructure and operational evidence sources. The important point is that we do not "
    "need to redesign the core system every time we add another source.")

# =========================================================
# SLIDE 8
# =========================================================
s = prs.slides.add_slide(BLANK); add_bg(s)
eyebrow(s, "What decision do we need now?")
title(s, "Validate the Production Path", size=32)

steps3 = ["Complete critical\nvalidation", "Run controlled\nproduction pilot", "Measure operational\nvalue"]
cw = Inches(3.75); gap = Inches(0.3); x = Inches(0.75); y = Inches(2.1)
for i, lbl in enumerate(steps3):
    card = rounded_card(s, x, y, cw, Inches(1.35), fill=WHITE, line=CARD_BORDER, radius=0.1)
    tf = card.text_frame; tf.word_wrap = True; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = Inches(0.25); tf.margin_right = Inches(0.2)
    p0 = tf.paragraphs[0]
    r0 = p0.add_run(); r0.text = str(i + 1)
    r0.font.size = Pt(26); r0.font.name = SERIF; r0.font.color.rgb = BLUE
    p1 = tf.add_paragraph()
    r1 = p1.add_run(); r1.text = lbl
    r1.font.size = Pt(13.5); r1.font.bold = True; r1.font.color.rgb = INK; r1.font.name = SANS
    x = Emu(int(x + cw + gap))

success = rounded_card(s, Inches(0.75), Inches(3.85), Inches(11.8), Inches(2.2), fill=AMBER_BG, line=AMBER_LINE, radius=0.08)
tf = success.text_frame; tf.margin_left = Inches(0.4); tf.margin_top = Inches(0.25)
tf.margin_right = Inches(0.4)
p = tf.paragraphs[0]
r = p.add_run(); r.text = "Success means more than a working demo."
r.font.size = Pt(17); r.font.italic = True; r.font.name = SERIF; r.font.color.rgb = INK

criteria = ["Accurate RCA", "Complete evidence", "Reliable execution", "Safe access controls", "Real incident validation"]
cy = Inches(4.55)
cx = Inches(1.15)
for c in criteria:
    dot = s.shapes.add_shape(MSO_SHAPE.OVAL, cx, cy + Inches(0.04), Inches(0.12), Inches(0.12))
    dot.fill.solid(); dot.fill.fore_color.rgb = AMBER
    dot.line.fill.background(); dot.shadow.inherit = False
    txt(s, Emu(int(cx + Inches(0.22))), cy, Inches(2.4), Inches(0.3), c, size=12.5, color=INK, font=SANS)
    cy = Emu(int(cy + Inches(0.42)))

set_notes(s,
    "My recommendation is not to jump directly from this demo into broad production use. The "
    "next step should be completing the remaining critical validation and then running a "
    "controlled pilot. We should measure whether the agent consistently gathers the right "
    "evidence, reaches accurate conclusions, and provides meaningful value to the engineer. If "
    "those results hold, we then have a strong basis for expanding the platform.")

out_path = "/private/tmp/claude-501/-Users-ashmin-projects/676bcf05-eba0-48a8-953d-6a5b8b22f991/scratchpad/SRE-Agent-Executive-Demo.pptx"
prs.save(out_path)
print("Saved:", out_path)
print("Slides:", len(prs.slides.__iter__.__self__._sldIdLst))
